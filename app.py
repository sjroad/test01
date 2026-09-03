# -*- coding: utf-8 -*-
"""
산지로드 결산 자동화 대시보드
------------------------------------------------
실행: streamlit run app.py

기능
1. 취합 결산서(xlsx) 업로드 -> 사이트별 규칙 적용 -> 최종 결산서 다운로드
2. 파이차트 대시보드 (사이트별 정산가/마진 비중, 상품 Top10)
3. 누적 데이터 검색 (사이트/상품/주문번호/날짜)
"""
import calendar
import io
from datetime import date
from pathlib import Path

import pandas as pd
import plotly.express as px
import sqlalchemy as sa
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import processing as P

DB_PATH = Path(__file__).parent / "data" / "settlement.db"

st.set_page_config(page_title="산지로드 결산 대시보드", layout="wide")


# ---------------------------------------------------------------------------
# 누적 DB
#
# 배포 환경(Streamlit Cloud 등)은 재시작/재배포 시 로컬 디스크가 초기화되므로,
# secrets에 [connections.sql] url이 설정돼 있으면 그 DB(예: Supabase Postgres)를
# 쓰고, 없으면(로컬 개발) 지금까지 쓰던 data/settlement.db 파일로 자동 대체한다.
# 팀 공유용으로 배포할 때는 반드시 secrets에 DB url을 등록해야 데이터가 유지된다.
# ---------------------------------------------------------------------------

@st.cache_resource
def get_engine() -> sa.Engine:
    try:
        has_cloud_db = "sql" in st.secrets.get("connections", {})
    except Exception:
        has_cloud_db = False  # secrets.toml 자체가 없는 로컬 개발 환경
    if has_cloud_db:
        return st.connection("sql", type="sql").engine
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    return sa.create_engine(f"sqlite:///{DB_PATH}")


def ensure_schema(engine: sa.Engine):
    with engine.begin() as conn:
        conn.execute(sa.text(
            """
            CREATE TABLE IF NOT EXISTS records (
                "결산일자" TEXT, "판매사" TEXT, "고객선택옵션" TEXT, "주문수량" REAL,
                "공급사배송비" REAL, "결제금액" REAL, "매입단가" REAL, "매입가" REAL,
                "정산가" REAL, "마진" REAL, "마진률" REAL, "적용규칙" TEXT,
                "비고" TEXT, "판매사주문번호" TEXT, "수령인연락처" TEXT
            )
            """
        ))


def save_to_db(df: pd.DataFrame, settlement_date: date):
    engine = get_engine()
    ensure_schema(engine)
    date_str = settlement_date.isoformat()
    with engine.begin() as conn:
        conn.execute(sa.text('DELETE FROM records WHERE "결산일자" = :d'), {"d": date_str})
    out = df.copy()
    out.insert(0, "결산일자", date_str)
    out.to_sql("records", engine, if_exists="append", index=False)


def load_all_from_db() -> pd.DataFrame:
    engine = get_engine()
    ensure_schema(engine)
    try:
        return pd.read_sql("SELECT * FROM records", engine)
    except Exception:
        return pd.DataFrame()


# ---------------------------------------------------------------------------
# 최종 결산서(xlsx) 생성
# ---------------------------------------------------------------------------

ACCOUNTING_FMT = '_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-'
DATE_HDR_FMT = r'mm"월" dd"일"'
_THIN = Side(style="thin")
THIN_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
GRAY_FILL = PatternFill("solid", fgColor="FFD3D3D3")
YELLOW_FILL = PatternFill("solid", fgColor="FFFFFF00")
BODY_FONT = Font(name="맑은 고딕", size=11)
TITLE_FONT = Font(name="맑은 고딕", size=18, bold=True)
DATE_FONT = Font(name="맑은 고딕", size=11, bold=True)
CENTER = Alignment(horizontal="center", vertical="center")


def _month_weekdays(year: int, month: int) -> list[date]:
    """해당 월의 평일(월~금) 목록. 공휴일은 반영하지 못한다(달력 정보 없음, 근사치)."""
    n_days = calendar.monthrange(year, month)[1]
    out = []
    for d in range(1, n_days + 1):
        cur = date(year, month, d)
        if cur.weekday() < 5:
            out.append(cur)
    return out


def _style_pivot_sheet(ws, n_data_rows: int, money_cols: list[str]):
    """'행 레이블 / 합계 : ...' 형태 시트 공통 서식(헤더·총합계 노란색, 테두리, 회계서식)."""
    header_row = 2
    last_row = header_row + n_data_rows  # 총합계 포함된 마지막 데이터 행
    for col in ["B"] + money_cols:
        c = ws[f"{col}{header_row}"]
        c.fill = YELLOW_FILL
        c.font = BODY_FONT
        c.border = THIN_BORDER
    for r in range(header_row + 1, last_row + 1):
        for col in ["B"] + money_cols:
            c = ws[f"{col}{r}"]
            c.font = BODY_FONT
            c.border = THIN_BORDER
            if col in money_cols:
                c.number_format = ACCOUNTING_FMT
        if r == last_row:  # 총합계 행
            for col in ["B"] + money_cols:
                ws[f"{col}{r}"].fill = YELLOW_FILL


def build_output_excel(df: pd.DataFrame, excluded: pd.DataFrame, settlement_date: date) -> bytes:
    """원본 '최종 결산서' 양식(열 너비·행 높이·병합 셀·색상·틀고정·자동필터 포함)과 최대한
    동일하게 만든다. 단 두 가지는 원본과 다르다:
    - 원본은 날짜별 매출 칸이 수식(SUM)으로 연결돼 있는데, 여기서는 그 시점에 계산된 값을
      그대로 채워 넣는다(매번 새로 생성되는 파일이라 하나로 이어지는 수식을 유지할 수 없음).
    - 날짜 트래커는 평일(월~금)만 나열한다. 추석 등 공휴일 달력 정보가 없어 반영하지 못했다.
    """
    site_summary = P.summarize_by_site(df)
    full_raw = pd.concat([df, excluded]).sort_index()

    sheet_df = full_raw.copy()
    for col in ["매입가", "정산가", "마진", "마진률"]:
        if col in sheet_df.columns:
            sheet_df[col] = sheet_df[col].fillna(0)
    sheet_df = sheet_df[[c for c in P.RAW_COLUMNS if c in sheet_df.columns]]

    product_all = (
        full_raw.groupby("고객선택옵션", dropna=False)["주문수량"].sum()
        .reset_index()
        .sort_values("주문수량", ascending=False)
        .rename(columns={"고객선택옵션": "행 레이블", "주문수량": "합계 : 주문수량"})
    )
    product_all = pd.concat([product_all, pd.DataFrame([{
        "행 레이블": "총합계", "합계 : 주문수량": product_all["합계 : 주문수량"].sum(),
    }])], ignore_index=True)

    site_tbl = site_summary[["판매사", "정산가", "마진"]].rename(
        columns={"판매사": "행 레이블", "정산가": "합계 : 정산가", "마진": "합계 : 마진"}
    )
    site_tbl = pd.concat([site_tbl, pd.DataFrame([{
        "행 레이블": "총합계",
        "합계 : 정산가": site_tbl["합계 : 정산가"].sum(),
        "합계 : 마진": site_tbl["합계 : 마진"].sum(),
    }])], ignore_index=True)

    # 당월 누적 = DB에 이미 저장된 이번 달 데이터(오늘 날짜 제외, 중복 방지) + 지금 이 화면의 오늘 데이터
    db_df = load_all_from_db()
    month_key = settlement_date.strftime("%Y-%m")
    date_str = settlement_date.isoformat()
    if not db_df.empty:
        prior_mask = (db_df["결산일자"].str.slice(0, 7) == month_key) & (db_df["결산일자"] != date_str)
        month_prior_rev = db_df.loc[prior_mask, "정산가"].sum()
        month_prior_margin = db_df.loc[prior_mask, "마진"].sum()
    else:
        month_prior_rev = month_prior_margin = 0.0
    today_rev = float(df["정산가"].sum())
    today_margin = float(df["마진"].sum())
    month_rev = month_prior_rev + today_rev
    month_margin = month_prior_margin + today_margin

    def day_totals(d: date):
        ds = d.isoformat()
        if ds == date_str:
            return today_rev, today_margin
        if not db_df.empty:
            hit = db_df[db_df["결산일자"] == ds]
            if not hit.empty:
                return float(hit["정산가"].sum()), float(hit["마진"].sum())
        return None, None

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        # ---------------- 결산서 ----------------
        sheet_df.to_excel(writer, sheet_name="결산서", index=False, startrow=8, startcol=1)
        ws = writer.sheets["결산서"]
        last_row = 9 + len(sheet_df)

        widths = {
            "A": 6.29, "B": 30.57, "C": 49.0, "D": 12.71, "E": 17.43, "F": 15.14,
            "G": 13.57, "H": 13.29, "I": 9.57, "J": 11.0, "K": 11.71, "L": 25.71,
            "M": 31.43, "N": 18.43, "O": 10.71, "P": 13.14, "Q": 14.43, "S": 14.29,
            "T": 13.14, "U": 14.43, "V": 13.14, "X": 14.43, "Y": 14.29, "Z": 14.43,
            "AA": 14.0, "AB": 13.57, "AC": 9.14,
        }
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        for r in [1, 3, 4, 5, 6, 7, 8, 9]:
            ws.row_dimensions[r].height = 16.5

        ws.freeze_panes = "A10"
        ws.auto_filter.ref = f"B9:AB{last_row}"

        ws.merge_cells("B1:L3")
        ws.merge_cells("B4:C8")
        ws.merge_cells("D4:E4"); ws.merge_cells("F4:G4")
        ws.merge_cells("D5:E5"); ws.merge_cells("F5:G5")
        ws.merge_cells("D6:E6"); ws.merge_cells("F6:G6")
        ws.merge_cells("D7:E7"); ws.merge_cells("F7:G7")
        ws.merge_cells("D8:L8")
        ws.merge_cells("H4:L7")

        ws["B1"] = f"{settlement_date.month:02d}월 {settlement_date.day:02d}일 산지로드 결산서"
        ws["B1"].font = TITLE_FONT
        ws["B1"].alignment = CENTER

        for label_cell, label_text, value_cell, value in [
            ("D4", "당월 누적 매출", "F4", month_rev),
            ("D5", "당월 누적 마진", "F5", month_margin),
            ("D6", "당일 매출", "F6", today_rev),
            ("D7", "당일 마진", "F7", today_margin),
        ]:
            lc = ws[label_cell]
            lc.value = label_text
            lc.font = BODY_FONT
            lc.fill = YELLOW_FILL
            lc.border = THIN_BORDER
            lc.alignment = CENTER
            vc = ws[value_cell]
            vc.value = value
            vc.font = BODY_FONT
            vc.number_format = ACCOUNTING_FMT
            vc.border = THIN_BORDER
            vc.alignment = CENTER

        # 날짜 트래커 (평일 12개씩 두 줄 — 원본의 P~AA 12칸 구성과 동일한 폭)
        weekdays = _month_weekdays(settlement_date.year, settlement_date.month)
        batch1, batch2 = weekdays[:12], weekdays[12:24]
        ws["O3"] = "당일매출"; ws["O3"].font = BODY_FONT
        ws["O4"] = "당일마진"; ws["O4"].font = BODY_FONT
        ws["O6"] = "당일매출"; ws["O6"].font = BODY_FONT
        ws["O7"] = "당일마진"; ws["O7"].font = BODY_FONT
        for batch, date_row, rev_row, margin_row in [(batch1, 2, 3, 4), (batch2, 5, 6, 7)]:
            for i, d in enumerate(batch):
                col = get_column_letter(16 + i)  # P부터
                dcell = ws[f"{col}{date_row}"]
                dcell.value = d
                dcell.number_format = DATE_HDR_FMT
                dcell.font = DATE_FONT
                dcell.fill = YELLOW_FILL
                rev, margin = day_totals(d)
                if rev is not None:
                    rc = ws[f"{col}{rev_row}"]
                    rc.value = rev
                    rc.number_format = ACCOUNTING_FMT
                    mc = ws[f"{col}{margin_row}"]
                    mc.value = margin
                    mc.number_format = ACCOUNTING_FMT

        for col in list("BCDEFGHIJKLMN"):
            c = ws[f"{col}9"]
            c.font = BODY_FONT
            c.fill = GRAY_FILL
            c.border = THIN_BORDER
        for r in range(10, last_row + 1):
            for col in list("BCDEFGHIJKLMN"):
                ws[f"{col}{r}"].border = THIN_BORDER
            for col in ["E", "F", "G", "H", "I", "J"]:
                ws[f"{col}{r}"].number_format = ACCOUNTING_FMT
            ws[f"K{r}"].number_format = "0%"

        # ---------------- 최종정산가 및 마진 ----------------
        site_tbl.to_excel(writer, sheet_name="최종정산가 및 마진", index=False, startrow=1, startcol=1)
        ws2 = writer.sheets["최종정산가 및 마진"]
        ws2.column_dimensions["B"].width = 30.57
        ws2.column_dimensions["C"].width = 15.43
        ws2.column_dimensions["D"].width = 14.29
        _style_pivot_sheet(ws2, len(site_tbl), ["C", "D"])

        # ---------------- 품목별 판매수량 ----------------
        product_all.to_excel(writer, sheet_name="품목별 판매수량", index=False, startrow=1, startcol=1)
        ws3 = writer.sheets["품목별 판매수량"]
        ws3.column_dimensions["B"].width = 70.29
        ws3.column_dimensions["C"].width = 16.71
        _style_pivot_sheet(ws3, len(product_all), ["C"])

    buf.seek(0)
    return buf.read()


def styled_site_table(site_summary: pd.DataFrame):
    """엑셀 피벗테이블 느낌(행 레이블 / 합계:정산가 / 합계:마진, 총합계 노란 강조)의 표."""
    tbl = site_summary[["판매사", "정산가", "마진"]].rename(
        columns={"판매사": "행 레이블", "정산가": "합계 : 정산가", "마진": "합계 : 마진"}
    )
    total_row = pd.DataFrame([{
        "행 레이블": "총합계",
        "합계 : 정산가": tbl["합계 : 정산가"].sum(),
        "합계 : 마진": tbl["합계 : 마진"].sum(),
    }])
    tbl = pd.concat([tbl, total_row], ignore_index=True)

    def highlight_total(row):
        if row["행 레이블"] == "총합계":
            return ["background-color: #FFFF00; font-weight: 700"] * len(row)
        return [""] * len(row)

    return (
        tbl.style.apply(highlight_total, axis=1)
        .format({"합계 : 정산가": "{:,.0f}", "합계 : 마진": "{:,.0f}"})
        .hide(axis="index")
    )


def kpi_box_html(rows: list[tuple[str, str]]) -> str:
    """이미지2의 노란색 라벨/값 박스 스타일 요약표."""
    trs = "".join(
        f'<tr><td style="background:#FFFF00;border:1px solid #999;padding:6px 14px;'
        f'font-weight:700;width:160px">{label}</td>'
        f'<td style="border:1px solid #999;padding:6px 14px;text-align:right;'
        f'min-width:160px">{value}</td></tr>'
        for label, value in rows
    )
    return f'<table style="border-collapse:collapse;font-size:15px">{trs}</table>'


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------

st.title("📊 산지로드 결산 자동화 대시보드")

tab_upload, tab_dashboard, tab_dashboard_monthly, tab_search, tab_rules = st.tabs(
    ["📥 결산서 업로드/처리", "🥧 대시보드", "📅 대시보드(월간누적)", "🔍 검색", "⚙️ 규칙 관리"]
)

# ---------------------------------------------------------------------------
# 탭 1: 업로드/처리
# ---------------------------------------------------------------------------
with tab_upload:
    st.subheader("1. 취합 결산서 업로드")
    settlement_date = st.date_input("결산 일자", value=date.today())
    uploaded = st.file_uploader("취합 결산서(xlsx) 파일을 올려주세요", type=["xlsx"])

    st.subheader("2. 토스 주문배송관리 파일 업로드")
    toss_uploaded = st.file_uploader(
        "🚚 토스 주문배송관리 파일 업로드 (선택 — 토스는 주문마다 수수료가 달라 "
        "이 파일이 있어야 정확하게 계산됩니다. 없으면 기본 96.7%로 계산됩니다)",
        type=["xlsx"],
        key="toss_uploader",
    )

    toss_rates = None
    if toss_uploaded is not None:
        try:
            toss_rates = P.read_toss_delivery_file(toss_uploaded)
            st.success(f"토스 배송관리 파일에서 주문 {len(toss_rates):,}건의 수수료 정보를 읽었습니다.")
        except Exception as e:
            st.error(f"토스 배송관리 파일을 읽는 중 오류가 발생했습니다: {e}")
            toss_rates = None

    raw = None
    result = None
    if uploaded is not None:
        try:
            raw = P.read_raw_settlement(uploaded)
        except Exception as e:
            st.error(f"파일을 읽는 중 오류가 발생했습니다: {e}")
            raw = None

        if raw is not None:
            st.success(f"{len(raw):,}건의 주문 데이터를 읽었습니다.")

            result = P.process_settlement(raw, toss_delivery_rates=toss_rates)

            if toss_rates is not None:
                toss_mask = result.df["판매사"].astype(str).str.contains("토스", na=False)
                toss_total = int(toss_mask.sum())
                toss_unmatched = int((toss_mask & result.df["정산가"].isna()).sum())
                if toss_total:
                    st.info(
                        f"토스 주문 {toss_total:,}건 중 {toss_total - toss_unmatched:,}건이 "
                        f"배송관리 파일과 매칭되었습니다 (미매칭 {toss_unmatched:,}건은 "
                        "'확인 필요' 목록에 포함됩니다 — 결산 대상 주문이 모두 담긴 배송관리 "
                        "파일인지 확인해 주세요)."
                    )
            n_total = len(result.df)
            n_ok = n_total - len(result.needs_review)

            c1, c2, c3, c4 = st.columns(4)
            c1.metric("전체 행", f"{n_total:,}")
            c2.metric("정산가 계산 완료", f"{n_ok:,}")
            c3.metric("확인 필요", f"{len(result.needs_review):,}")
            c4.metric("결산 제외", f"{len(result.excluded):,}")

            if len(result.excluded) > 0:
                with st.expander(
                    f"결산 대상 18개 사이트에 없어 제외한 행 {len(result.excluded):,}건 "
                    "(맛장군/쿠팡 등)"
                ):
                    st.dataframe(
                        result.excluded[["판매사", "고객선택옵션", "결제금액", "주문수량"]]
                        .groupby("판매사", as_index=False)
                        .agg(건수=("결제금액", "count"), 결제금액합계=("결제금액", "sum"))
                        .sort_values("결제금액합계", ascending=False),
                        use_container_width=True,
                    )

            if len(result.needs_review) > 0:
                st.warning(
                    "아래 사이트/상품은 정산 규칙이 없어 정산가를 계산하지 못했습니다. "
                    "'규칙 관리' 탭에서 요율 또는 고정 정산단가를 등록한 뒤 다시 처리해 주세요."
                )
                st.dataframe(
                    result.needs_review[["판매사", "고객선택옵션", "결제금액", "주문수량"]]
                    .groupby(["판매사", "고객선택옵션"], as_index=False)
                    .agg(건수=("결제금액", "count"), 결제금액합계=("결제금액", "sum")),
                    use_container_width=True,
                )

            st.caption("계산 결과")
            st.dataframe(result.df, use_container_width=True, height=350)
            st.session_state["last_result_df"] = result.df

    # -----------------------------------------------------------------
    # 3. 결산서 최종본 다운로드 — 1번 업로드 전에도 버튼은 항상 보이되,
    #    데이터가 없으면 비활성화 상태로 둔다.
    # -----------------------------------------------------------------
    st.subheader("3. 결산서 최종본 다운로드")
    col_a, col_b = st.columns(2)
    with col_a:
        if st.button("💾 누적 DB에 저장 (이 날짜로 반영)", type="primary", disabled=result is None):
            save_to_db(result.df, settlement_date)
            st.success(f"{settlement_date} 데이터를 누적 DB에 저장했습니다.")
            st.session_state["_reload_db"] = True
    with col_b:
        st.download_button(
            "⬇️ 결산서 최종본 다운로드 (.xlsx)",
            data=build_output_excel(result.df, result.excluded, settlement_date) if result is not None else b"",
            file_name=f"{settlement_date.strftime('%m%d')}_산지로드_결산_최종.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            disabled=result is None,
        )
    if result is None:
        st.caption("먼저 1번에서 취합 결산서를 업로드하면 활성화됩니다.")

    # -----------------------------------------------------------------
    # 4. 담당자별 수정/검토 파일 업로드
    # -----------------------------------------------------------------
    st.subheader("4. 담당자별 수정/검토 파일 업로드")
    st.caption(
        "위 3번에서 받은 '결산서 최종본'을 담당자가 열어 담당 사이트의 매입단가(원가)만 "
        "확인/수정한 뒤 그대로 다시 올려주세요. 본인 담당이 아닌 사이트 행은 값이 "
        "바뀌어 있어도 전부 무시되니, 실수로 다른 사이트를 건드려도 반영되지 않습니다."
    )

    working_df = None
    if result is not None:
        base_sig = f"{settlement_date.isoformat()}::{len(raw)}"
        if st.session_state.get("reviewed_base_sig") != base_sig:
            st.session_state["reviewed_df"] = result.df.copy()
            st.session_state["reviewed_base_sig"] = base_sig
        working_df = st.session_state["reviewed_df"]

    reviewer_cols = st.columns(len(P.REVIEWER_SITE_KEYWORDS))
    for col, (name, keywords) in zip(reviewer_cols, P.REVIEWER_SITE_KEYWORDS.items()):
        with col:
            st.markdown(f"**{name}**")
            st.caption(", ".join(keywords))
            rfile = st.file_uploader(
                f"{name} 검토 파일", type=["xlsx"], key=f"reviewer_{name}",
                disabled=result is None,
            )
            if rfile is not None and working_df is not None:
                try:
                    reviewed_raw = P.read_raw_settlement(rfile)
                except Exception as e:
                    st.error(f"파일을 읽는 중 오류가 발생했습니다: {e}")
                    reviewed_raw = None
                if reviewed_raw is not None:
                    working_df, stats = P.apply_reviewer_corrections(
                        working_df, reviewed_raw, keywords
                    )
                    st.success(
                        f"담당 {stats['owned_rows']:,}행 중 매입단가 "
                        f"{stats['changed_rows']:,}건 반영"
                    )
                    if stats["ignored_rows"] > 0:
                        st.caption(
                            f"※ 다른 담당자 사이트 행 {stats['ignored_rows']:,}건은 무시했습니다."
                        )

    if result is None:
        st.caption("먼저 1번에서 취합 결산서를 업로드하면 검토 파일을 올릴 수 있습니다.")
    else:
        st.session_state["reviewed_df"] = working_df
        unassigned = P.unassigned_sites(result.df)
        if unassigned:
            st.caption(
                "⚠ 세 담당자 누구에게도 배정되지 않아 검토 대상에서 빠진 사이트: "
                + ", ".join(unassigned)
            )

    # -----------------------------------------------------------------
    # 5. 담당자 검토 반영 후 결산서 최종본 다운로드
    # -----------------------------------------------------------------
    st.subheader("5. 결산서 최종본 다운로드 (담당자 검토 반영)")
    col_c, col_d = st.columns(2)
    with col_c:
        if st.button("💾 누적 DB에 저장 (검토 반영본)", key="save_reviewed", disabled=working_df is None):
            save_to_db(working_df, settlement_date)
            st.success(f"{settlement_date} 데이터를(검토 반영본) 누적 DB에 저장했습니다.")
            st.session_state["_reload_db"] = True
    with col_d:
        st.download_button(
            "⬇️ 결산서 최종본 다운로드 (검토 반영, .xlsx)",
            data=build_output_excel(working_df, result.excluded, settlement_date) if working_df is not None else b"",
            file_name=f"{settlement_date.strftime('%m%d')}_산지로드_결산_최종_검토반영.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="download_reviewed",
            disabled=working_df is None,
        )
    if working_df is None:
        st.caption("먼저 1번에서 취합 결산서를 업로드하면 활성화됩니다.")

# ---------------------------------------------------------------------------
# 탭 2: 대시보드
# ---------------------------------------------------------------------------
with tab_dashboard:
    st.subheader("누적 대시보드")
    db_df = load_all_from_db()

    if db_df.empty:
        st.info("아직 저장된 누적 데이터가 없습니다. '결산서 업로드/처리' 탭에서 먼저 저장해 주세요.")
    else:
        dates = sorted(db_df["결산일자"].unique())
        sel_dates = st.multiselect("기간(결산일자) 선택 - 비워두면 전체", dates, default=dates)
        view_df = db_df[db_df["결산일자"].isin(sel_dates)] if sel_dates else db_df

        total_i = view_df["정산가"].sum()
        total_j = view_df["마진"].sum()
        c1, c2, c3 = st.columns(3)
        c1.metric("합계 정산가", f"{total_i:,.0f}")
        c2.metric("합계 마진", f"{total_j:,.0f}")
        c3.metric("마진률", f"{(total_j/total_i):.1%}" if total_i else "-")

        site_summary = P.summarize_by_site(view_df)
        product_summary = P.summarize_by_product(view_df)

        st.markdown("**사이트별 정산가 · 마진 합계**")
        st.dataframe(styled_site_table(site_summary), use_container_width=True, hide_index=True)

        col1, col2 = st.columns(2)
        with col1:
            fig1 = px.pie(
                site_summary.dropna(subset=["정산가"]),
                names="판매사", values="정산가",
                title="사이트별 정산가 비중",
            )
            st.plotly_chart(fig1, use_container_width=True)

        with col2:
            fig2 = px.pie(
                site_summary.dropna(subset=["마진"]),
                names="판매사", values="마진",
                title="사이트별 마진 비중",
            )
            st.plotly_chart(fig2, use_container_width=True)

        top10 = product_summary.head(10).copy()
        etc_qty = product_summary["주문수량"].iloc[10:].sum()
        if etc_qty > 0:
            top10 = pd.concat(
                [top10, pd.DataFrame([{"고객선택옵션": "기타", "주문수량": etc_qty}])],
                ignore_index=True,
            )
        fig3 = px.pie(top10, names="고객선택옵션", values="주문수량",
                      title="상품별 주문수량 비중 (Top10 + 기타)")
        st.plotly_chart(fig3, use_container_width=True)

        st.subheader("월별 추이")
        monthly = view_df.copy()
        monthly["월"] = monthly["결산일자"].str.slice(0, 7)
        monthly_summary = monthly.groupby("월").agg(
            결제금액=("결제금액", "sum"), 정산가=("정산가", "sum"), 마진=("마진", "sum")
        ).reset_index()
        fig4 = px.bar(monthly_summary, x="월", y=["결제금액", "정산가", "마진"], barmode="group",
                      title="월별 결제금액/정산가/마진")
        st.plotly_chart(fig4, use_container_width=True)

# ---------------------------------------------------------------------------
# 탭 2-2: 대시보드(월간누적)
# ---------------------------------------------------------------------------
with tab_dashboard_monthly:
    st.subheader("월간 누적 대시보드")
    db_df = load_all_from_db()

    if db_df.empty:
        st.info("아직 저장된 누적 데이터가 없습니다. '결산서 업로드/처리' 탭에서 먼저 저장해 주세요.")
    else:
        months = sorted(db_df["결산일자"].str.slice(0, 7).unique(), reverse=True)
        sel_month = st.selectbox("월 선택", months)
        month_df = db_df[db_df["결산일자"].str.slice(0, 7) == sel_month]

        month_total = month_df["정산가"].sum()
        month_margin = month_df["마진"].sum()
        st.markdown(
            kpi_box_html([
                (f"{sel_month} 누적 매출", f"{month_total:,.0f}"),
                (f"{sel_month} 누적 마진", f"{month_margin:,.0f}"),
            ]),
            unsafe_allow_html=True,
        )

        st.write("")
        month_site_summary = P.summarize_by_site(month_df)
        col1, col2 = st.columns(2)
        with col1:
            fig_m1 = px.pie(
                month_site_summary.dropna(subset=["정산가"]),
                names="판매사", values="정산가",
                title=f"{sel_month} 사이트별 누적 매출 비중",
            )
            st.plotly_chart(fig_m1, use_container_width=True)
        with col2:
            fig_m2 = px.pie(
                month_site_summary.dropna(subset=["마진"]),
                names="판매사", values="마진",
                title=f"{sel_month} 사이트별 누적 마진 비중",
            )
            st.plotly_chart(fig_m2, use_container_width=True)

# ---------------------------------------------------------------------------
# 탭 3: 검색
# ---------------------------------------------------------------------------
with tab_search:
    st.subheader("누적 데이터 검색")
    db_df = load_all_from_db()

    if db_df.empty:
        st.info("아직 저장된 누적 데이터가 없습니다.")
    else:
        c1, c2, c3 = st.columns(3)
        site_q = c1.text_input("판매사(부분검색)")
        product_q = c2.text_input("상품명(부분검색)")
        order_q = c3.text_input("주문번호(부분검색)")

        all_dates = sorted(db_df["결산일자"].unique())
        min_d = date.fromisoformat(all_dates[0])
        max_d = date.fromisoformat(all_dates[-1])
        date_range = st.date_input(
            "결산일자 기간 (기간 ~ 기간)", value=(min_d, max_d), min_value=min_d, max_value=max_d,
        )
        if isinstance(date_range, tuple) and len(date_range) == 2:
            start_d, end_d = date_range
        else:
            start_d = end_d = date_range[0] if isinstance(date_range, tuple) else date_range

        filtered = db_df.copy()
        if site_q:
            filtered = filtered[filtered["판매사"].str.contains(site_q, case=False, na=False)]
        if product_q:
            filtered = filtered[filtered["고객선택옵션"].str.contains(product_q, case=False, na=False)]
        if order_q:
            filtered = filtered[filtered["판매사주문번호"].astype(str).str.contains(order_q, case=False, na=False)]
        filtered = filtered[
            (filtered["결산일자"] >= start_d.isoformat()) & (filtered["결산일자"] <= end_d.isoformat())
        ]

        st.write(f"검색 결과: {len(filtered):,}건")
        st.dataframe(filtered, use_container_width=True, height=450)

        if len(filtered) > 0:
            t1, t2 = st.columns(2)
            t1.metric("매출(정산가) 합계", f"{filtered['정산가'].sum():,.0f}")
            t2.metric("마진 합계", f"{filtered['마진'].sum():,.0f}")

            csv_bytes = filtered.to_csv(index=False).encode("utf-8-sig")
            st.download_button("⬇️ 검색 결과 CSV 다운로드", data=csv_bytes, file_name="검색결과.csv")

# ---------------------------------------------------------------------------
# 탭 4: 규칙 관리
# ---------------------------------------------------------------------------
with tab_rules:
    st.subheader("사이트별 정산 규칙")
    st.caption(
        "판매사명에 '키워드'가 포함되면 해당 규칙이 적용됩니다. "
        "type: percent(결제금액×rate) / same_as_payment(결제금액과 동일, rate=1.0)"
    )
    rules = P.load_rate_rules()
    rules_df = pd.DataFrame(rules)
    edited_rules = st.data_editor(rules_df, num_rows="dynamic", use_container_width=True, key="rules_editor")
    if st.button("규칙 저장", key="save_rules"):
        P.save_rate_rules(edited_rules.to_dict("records"))
        st.success("사이트 규칙을 저장했습니다.")

    st.divider()
    st.subheader("캐시딜 예외 상품")
    ce = P.load_cashdeal_exceptions()
    ce_df = pd.DataFrame(ce)
    edited_ce = st.data_editor(ce_df, num_rows="dynamic", use_container_width=True, key="ce_editor")
    if st.button("캐시딜 예외 저장", key="save_ce"):
        P.save_cashdeal_exceptions(edited_ce.to_dict("records"))
        st.success("캐시딜 예외 상품을 저장했습니다.")

    st.divider()
    st.subheader("고정 정산단가 상품 (케이딜 / 꿈꾸는이웃 / 홈앤쇼핑 / LG복지몰 / 제트언스)")
    st.caption(
        "이 사이트들은 결제금액이 아니라 상품별로 정해진 '개당 정산단가'를 사용합니다. "
        "정산가 = 정산단가 × 주문수량. 새 상품이 나오면 이 표에 추가해 주세요."
    )
    fp = P.load_fixed_price_table()
    edited_fp = st.data_editor(fp, num_rows="dynamic", use_container_width=True, key="fp_editor")
    if st.button("고정 정산단가 저장", key="save_fp"):
        P.save_fixed_price_table(edited_fp)
        st.success("고정 정산단가 표를 저장했습니다.")
